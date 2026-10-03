"""Evaluate the Python rule baseline against fixed development labels.

Run this file after ``python app/run_baseline.py``:
    python app/evaluate_baseline.py

It compares ``data/baseline_output_development.xlsx`` with
``data/ground_truth_development.xlsx`` and writes
``data/baseline_evaluation_development.xlsx``. No API key or model call is used.
"""

from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
GROUND_TRUTH_PATH = DATA_DIR / "ground_truth_development.xlsx"
BASELINE_PATH = DATA_DIR / "baseline_output_development.xlsx"
OUTPUT_PATH = DATA_DIR / "baseline_evaluation_development.xlsx"

TRACKED_FIELDS = [
    ("Category", "Category"),
    ("Price USD", "Price USD"),
    ("Condition", "Condition"),
    ("Transaction type", "Transaction type"),
    ("Needs human review", "Needs human review"),
]


def same_value(left, right):
    """Compare values case-insensitively, while allowing equivalent numerics."""
    if isinstance(left, (int, float)) and not isinstance(left, bool) and isinstance(right, (int, float)) and not isinstance(right, bool):
        return abs(left - right) < 0.000001
    return str(left).strip().lower() == str(right).strip().lower()


def style_sheet(sheet, widths, header_center=True):
    """Apply plain black text, white fills, wrapping, and visible grey borders."""
    thin_grey = Side(style="thin", color="D9D9D9")
    header_font = Font(name="Arial", size=10, bold=True, color="000000")
    body_font = Font(name="Arial", size=10, color="000000")
    white = PatternFill("solid", fgColor="FFFFFF")
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
            cell.fill = white
            cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
            cell.alignment = Alignment(
                horizontal="center" if cell.row == 1 and header_center else "left",
                vertical="center",
                wrap_text=True,
            )
    for row in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 28


def main():
    """Join fixed labels to predictions, calculate accuracy, and save the workbook."""
    truth = pd.read_excel(GROUND_TRUTH_PATH, sheet_name="Ground truth")
    predicted = pd.read_excel(BASELINE_PATH, sheet_name="Baseline output")
    predicted_by_id = predicted.set_index("Post ID").to_dict(orient="index")

    comparison_rows = []
    for _, expected in truth.iterrows():
        post_id = expected["Post ID"]
        if post_id not in predicted_by_id:
            raise ValueError(f"Missing baseline prediction for {post_id}")
        prediction = predicted_by_id[post_id]
        matches = [same_value(expected[truth_column], prediction[prediction_column]) for truth_column, prediction_column in TRACKED_FIELDS]
        comparison_rows.append(
            {
                "Post ID": post_id,
                "Category truth": expected["Category"],
                "Category baseline": prediction["Category"],
                "Category correct": matches[0],
                "Price truth": expected["Price USD"],
                "Price baseline": prediction["Price USD"],
                "Price correct": matches[1],
                "Condition truth": expected["Condition"],
                "Condition baseline": prediction["Condition"],
                "Condition correct": matches[2],
                "Transaction truth": expected["Transaction type"],
                "Transaction baseline": prediction["Transaction type"],
                "Transaction correct": matches[3],
                "Review truth": expected["Needs human review"],
                "Review baseline": prediction["Needs human review"],
                "Review correct": matches[4],
                "All tracked fields correct": all(matches),
            }
        )

    comparison = pd.DataFrame(comparison_rows)
    metric_rows = [["Cases evaluated", len(comparison), len(comparison), 1.0]]
    field_to_column = [
        ("Category accuracy", "Category correct"),
        ("Price USD accuracy", "Price correct"),
        ("Condition accuracy", "Condition correct"),
        ("Transaction type accuracy", "Transaction correct"),
        ("Needs human review accuracy", "Review correct"),
        ("All tracked fields correct", "All tracked fields correct"),
    ]
    for metric, column in field_to_column:
        correct = int(comparison[column].sum())
        metric_rows.append([metric, correct, len(comparison), correct / len(comparison)])
    metrics = pd.DataFrame(metric_rows, columns=["Metric", "Correct cases", "Cases evaluated", "Accuracy"])
    method = pd.DataFrame(
        [
            ["Data", "40 development posts in trade_posts.xlsx."],
            ["Truth source", "Fixed labels in ground_truth_development.xlsx, created before the baseline was evaluated."],
            ["Prediction source", "Keyword outputs in baseline_output_development.xlsx."],
            ["Scoring", "Exact-match accuracy for category, price USD, condition, transaction type, and needs-human-review. Product name is excluded because the rule baseline does not extract names."],
            ["Interpretation", "This is development-set performance. It guides AI prompt design and is not the final holdout result."],
        ],
        columns=["Item", "Explanation"],
    )

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        metrics.to_excel(writer, sheet_name="Metrics", index=False)
        comparison.to_excel(writer, sheet_name="Post comparison", index=False)
        method.to_excel(writer, sheet_name="Method", index=False)
        style_sheet(writer.book["Metrics"], {"A": 30, "B": 18, "C": 18, "D": 14})
        style_sheet(writer.book["Post comparison"], {"A": 12, **{chr(66 + index): 18 for index in range(16)}})
        style_sheet(writer.book["Method"], {"A": 18, "B": 100})
        writer.book["Post comparison"].freeze_panes = "A2"
        for cell in writer.book["Metrics"]["D"][1:]:
            cell.number_format = "0.0%"

    print(metrics.to_string(index=False))
    print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
