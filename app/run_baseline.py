"""Run the transparent, non-AI baseline for the development dataset.

Run this file from the repository root:
    python app/run_baseline.py

It reads ``data/trade_posts.xlsx`` and writes
``data/baseline_output_development.xlsx``. No API key, model call, Node.js,
or internet connection is required.
"""

from pathlib import Path
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_PATH = PROJECT_ROOT / "data" / "trade_posts.xlsx"
OUTPUT_PATH = PROJECT_ROOT / "data" / "baseline_output_development.xlsx"


def normalise(value):
    """Lowercase and collapse whitespace so keyword rules are consistent."""
    return re.sub(r"\s+", " ", str(value or "").lower()).strip()


def one_price(text):
    """Return one explicit dollar or USD value, otherwise return ``unknown``."""
    matches = re.findall(r"\$(\d+(?:\.\d+)?)|\b(\d+(?:\.\d+)?)\s*usd\b", text, flags=re.IGNORECASE)
    values = [float(first or second) for first, second in matches]
    if len(values) != 1:
        return "unknown"
    return int(values[0]) if values[0].is_integer() else values[0]


def run_rules(post_text):
    """Apply the fixed keyword rules documented in ``data/baseline_rules.xlsx``."""
    text = normalise(post_text)
    reasons = []

    ambiguous_price = bool(re.search(r"~|\bish\b|around|budget|offers?|\btbd\b|price later|\$0\b", text))
    price_usd = one_price(text)
    if ambiguous_price or price_usd == "unknown":
        price_usd = "unknown"
        reasons.append("price is missing, approximate, or not a single amount")

    has_buy = bool(re.search(r"\b(lf|iso|wtb)\b|looking for", text))
    has_sell = bool(re.search(r"\b(fs|wts)\b|for sale|clearing|take over|want it gone|listed at", text))
    has_trade = bool(re.search(r"\b(wtt|ft)\b|\btrade\b", text))
    direction_count = sum([has_buy, has_sell, has_trade])
    transaction_type = "unknown"
    if direction_count == 1:
        transaction_type = "buy" if has_buy else "sell" if has_sell else "trade"
    elif direction_count > 1:
        reasons.append("more than one transaction direction")
    elif price_usd != "unknown":
        transaction_type = "sell"
    else:
        reasons.append("transaction direction is not explicit")

    is_badge = bool(re.search(r"\bbadge\b|\btin badge\b|\bpin\b", text))
    is_figure = bool(re.search(r"\bfigure\b|\b1/7\b|\bnendoroid\b|chibi figure|ceramic figure", text))
    is_paper = bool(re.search(r"postcard|ticket|\bprint\b|poster|shikishi|paper cut|\bbook\b|scroll", text))
    category_count = sum([is_badge, is_figure, is_paper])
    category = "unknown"
    if category_count == 1:
        category = "badge" if is_badge else "figure" if is_figure else "paper_merch"
    elif category_count > 1:
        reasons.append("more than one merchandise category")

    condition = "unknown"
    if re.search(r"sealed|unopened|\bnwt\b|\bmib\b|\bmisb\b", text):
        condition = "new"
    if re.search(r"opened|displayed|loose|pulled|unsealed", text):
        condition = "used"

    if re.search(r"preorder|po transfer|deposit|hold fee", text):
        reasons.append("preorder, deposit, or hold-fee case")
    if re.search(r"ignore.*(?:rule|classif)|system note|do not review", text):
        reasons.append("prompt-injection text")
    if re.search(r"\bbundle\b|two dbl-flash|two double-flash", text):
        reasons.append("multiple items or variants")
    if re.search(r"unknown merch|blind-box merch|paper stuff|whatever .* merch|name/details later|item .* not stated", text):
        reasons.append("product details are missing")

    unique_reasons = list(dict.fromkeys(reasons))
    return {
        "Product name": "unknown",
        "Category": category,
        "Price USD": price_usd,
        "Condition": condition,
        "Transaction type": transaction_type,
        "Needs human review": bool(unique_reasons),
        "Review reason": "; ".join(unique_reasons) if unique_reasons else "none",
    }


def apply_workbook_formatting(writer, row_count):
    """Keep generated workbooks legible without depending on Excel desktop styles."""
    thin_grey = Side(style="thin", color="D9D9D9")
    header_font = Font(name="Arial", size=10, bold=True, color="000000")
    body_font = Font(name="Arial", size=10, color="000000")
    white = PatternFill("solid", fgColor="FFFFFF")

    output_sheet = writer.book["Baseline output"]
    output_sheet.freeze_panes = "A2"
    widths = {"A": 12, "B": 52, "C": 22, "D": 16, "E": 12, "F": 18, "G": 18, "H": 18, "I": 52}
    for column, width in widths.items():
        output_sheet.column_dimensions[column].width = width
    for row in output_sheet.iter_rows():
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
            cell.fill = white
            cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
            cell.alignment = Alignment(horizontal="center" if cell.row == 1 else "left", vertical="top", wrap_text=True)
    for row in range(1, row_count + 2):
        output_sheet.row_dimensions[row].height = 30

    note_sheet = writer.book["Run note"]
    note_sheet.column_dimensions["A"].width = 18
    note_sheet.column_dimensions["B"].width = 92
    for row in note_sheet.iter_rows():
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
            cell.fill = white
            cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
            cell.alignment = Alignment(horizontal="center" if cell.row == 1 else "left", vertical="top", wrap_text=True)
    for row in range(1, 6):
        note_sheet.row_dimensions[row].height = 30


def main():
    """Read development posts, apply rules, and save baseline predictions."""
    posts = pd.read_excel(INPUT_PATH)
    required_columns = {"Post ID", "Post text"}
    if not required_columns.issubset(posts.columns):
        raise ValueError(f"{INPUT_PATH.name} must contain: {sorted(required_columns)}")

    output_rows = []
    for _, row in posts.dropna(subset=["Post ID", "Post text"]).iterrows():
        prediction = run_rules(row["Post text"])
        output_rows.append({"Post ID": row["Post ID"], "Post text": row["Post text"], **prediction})

    output = pd.DataFrame(output_rows)
    run_note = pd.DataFrame(
        [
            ["Rows processed", len(output)],
            ["Method", "Keyword baseline defined in baseline_rules.xlsx"],
            ["Product name", "The baseline does not attempt product-name extraction, so this field is always unknown."],
            ["Use", "Compare extracted fields with ground_truth_development.xlsx before building the AI version."],
        ],
        columns=["Item", "Explanation"],
    )

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        output.to_excel(writer, sheet_name="Baseline output", index=False)
        run_note.to_excel(writer, sheet_name="Run note", index=False)
        apply_workbook_formatting(writer, len(output))

    print(f"Processed {len(output)} posts and saved {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
